import io
import os
import re
import sys

import mongomock
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app  # noqa: E402
from config import Settings  # noqa: E402

GOOD_RESULT = {
    "match_score": 72,
    "summary": "Solid backend base.",
    "skills": ["Python", "Flask"],
    "missing_skills": ["Docker"],
    "roadmap": ["Learn Docker basics"],
    "interview_questions": ["Explain the GIL."],
    "model": "fake-model",
}
RESUME = "Python developer with 3 years of experience building Flask APIs and MongoDB backends. " * 2


class FakeAnalyzer:
    def __init__(self, result=None):
        self.result = GOOD_RESULT if result is None else result
        self.calls = []

    def __call__(self, resume_text, role):
        self.calls.append((resume_text, role))
        return dict(self.result)


@pytest.fixture
def settings():
    return Settings(secret_key="x" * 40, mongodb_uri="mongodb://localhost", max_upload_mb=1)


@pytest.fixture
def database():
    return mongomock.MongoClient().career_test


@pytest.fixture
def analyzer():
    return FakeAnalyzer()


@pytest.fixture
def app(settings, database, analyzer):
    return create_app(settings, database=database, analyzer=analyzer)


@pytest.fixture
def client(app):
    return app.test_client()


def get_token(client, path="/login"):
    html = client.get(path).get_data(as_text=True)
    return re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)


def post(client, path, data=None, token_from="/login", **kw):
    data = dict(data or {})
    data["csrf_token"] = get_token(client, token_from)
    return client.post(path, data=data, **kw)


def signup_and_login(client, email="a@example.com", password="password123"):
    post(client, "/signup", {"email": email, "password": password}, token_from="/signup")
    return post(client, "/login", {"email": email, "password": password})


def make_pdf(text):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for line in text.split("\n"):
        c.drawString(50, y, line)
        y -= 18
    c.save()
    return buf.getvalue()


def make_docx(text, table_cell=None):
    import docx
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    if table_cell:
        t = d.add_table(rows=1, cols=1)
        t.rows[0].cells[0].text = table_cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
