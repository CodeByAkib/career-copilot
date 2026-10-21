"""Extract text from uploaded PDF/DOCX resumes. Files are processed in memory only."""
import io
import zipfile

import docx
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from validation import clean_text

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_PDF_PAGES = 30
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # guards against zip bombs


class ResumeInputError(ValueError):
    """A user-facing problem with the uploaded resume file."""


def _extension(filename: str) -> str:
    name = (filename or "").lower().strip()
    return name[name.rfind("."):] if "." in name else ""


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ResumeInputError("This PDF is password-protected. Please upload an unlocked copy.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ResumeInputError(f"This PDF has more than {MAX_PDF_PAGES} pages; please upload just your resume.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ResumeInputError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError, TypeError):
        raise ResumeInputError("Unable to read this PDF. Please upload a valid, uncorrupted PDF.") from None


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ResumeInputError("This file is not a valid DOCX document.")
            if sum(i.file_size for i in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise ResumeInputError("This DOCX file is too large once unpacked.")
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:  # many resume templates use tables
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except ResumeInputError:
        raise
    except Exception:
        raise ResumeInputError("Unable to read this DOCX. Please upload a valid, uncorrupted DOCX.") from None


def extract_resume_text(filename: str, data: bytes) -> str:
    """Return cleaned text from a PDF/DOCX. The content is verified, not just the filename."""
    ext = _extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeInputError("Unsupported file type. Please upload a PDF or DOCX file.")
    if not data:
        raise ResumeInputError("The uploaded file is empty.")

    if ext == ".pdf":
        if not data.lstrip()[:5] == b"%PDF-":
            raise ResumeInputError("This file is not a real PDF, even though its name ends in .pdf.")
        text = _extract_pdf(data)
    else:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ResumeInputError("This file is not a real DOCX, even though its name ends in .docx.")
        text = _extract_docx(data)

    text = clean_text(text)
    if not text:
        raise ResumeInputError(
            "No text could be extracted from this file. Scanned/image-only resumes are not "
            "supported; please upload a text-based file or paste your resume."
        )
    return text
